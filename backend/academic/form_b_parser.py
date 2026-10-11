import csv
import io
import re
from pathlib import Path

COURSE_CODE_RE = re.compile(r"\b([A-Z]{2,8}[\s-]?\d{3,5}[A-Z]?)\b", re.I)
CREDIT_STATUS_RE = re.compile(r"\b\d{1,2}\s+[A-Z]\s*$", re.I)


def _clean(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


def _dedupe(items):
    out = []
    seen = set()
    for code, name in items:
        code = re.sub(r"[\s-]+", "", code).upper()
        name = _clean(name).strip("-–—|:")
        if not name or len(name) < 3:
            continue
        key = code.lower()
        if key not in seen:
            seen.add(key)
            out.append({"code": code, "name": name})
    return out


def _extract_table_rows(rows):
    found = []
    for row in rows:
        cells = [_clean(x) for x in row if _clean(x)]
        if not cells:
            continue
        text = " | ".join(cells)
        match = COURSE_CODE_RE.search(text)
        if not match:
            continue
        code = match.group(1)
        # Prefer text after the code in the same cell/row; otherwise the next cell.
        after = text[match.end():]
        after = re.sub(r"^\s*[-:|]\s*", "", after)
        name = after
        if not name:
            code_cell = next((i for i, c in enumerate(cells) if COURSE_CODE_RE.search(c)), None)
            if code_cell is not None and code_cell + 1 < len(cells):
                name = cells[code_cell + 1]
        if name:
            name = re.sub(r"\b\d{1,2}\s+[A-Z]\s*$", "", name, flags=re.I)
            found.append((code, name))
    return _dedupe(found)


def _extract_pdf(text):
    lines = [_clean(line) for line in text.splitlines()]
    found = []
    i = 0
    while i < len(lines):
        line = lines[i]
        # pypdf commonly puts the S/N and course code on separate lines.
        match = COURSE_CODE_RE.fullmatch(line)
        if not match:
            i += 1
            continue
        code = match.group(1)
        chunks = []
        j = i + 1
        while j < len(lines):
            nxt = lines[j]
            if not nxt:
                j += 1
                continue
            # Credit value is immediately followed by course status (C/E).
            if re.fullmatch(r"\d{1,2}", nxt) and j + 1 < len(lines) and re.fullmatch(r"[A-Za-z]", lines[j + 1]):
                break
            # A new course code starts the next record.
            if COURSE_CODE_RE.fullmatch(nxt):
                break
            # Section/header and footer boundaries.
            if re.match(r"^(Major Courses|Elective Courses|Total Credit Value|Course Instructor|Course Status)$", nxt, re.I):
                break
            chunks.append(nxt)
            j += 1
        name = _clean(" ".join(chunks))
        if name and len(name) > 2:
            found.append((code, name))
        i = max(j, i + 1)
    return _dedupe(found)


def parse_form_b(file_obj, filename=None):
    filename = (filename or getattr(file_obj, "name", "") or "").lower()
    data = file_obj.read()
    if hasattr(file_obj, "seek"):
        file_obj.seek(0)
    suffix = Path(filename).suffix
    rows = []
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows.extend(ws.iter_rows(values_only=True))
        courses = _extract_table_rows(rows)
    elif suffix == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="ignore"))))
        courses = _extract_table_rows(rows)
    elif suffix in {".txt", ".text"}:
        rows = [re.split(r"\t|,|;|\s{2,}", line) for line in data.decode("utf-8-sig", errors="ignore").splitlines()]
        courses = _extract_table_rows(rows)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        rows.extend([[c.text for c in table.rows[r].cells] for table in doc.tables for r in range(len(table.rows))])
        rows.extend([[p.text] for p in doc.paragraphs])
        courses = _extract_table_rows(rows)
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        courses = []
        for page in reader.pages:
            courses.extend(_extract_pdf(page.extract_text() or ""))
        courses = _dedupe([(x["code"], x["name"]) for x in courses])
    else:
        raise ValueError("Unsupported Form B file. Use PDF, Excel (.xlsx/.xlsm), Word (.docx), CSV or TXT.")
    if not courses:
        raise ValueError("No course codes were detected in the Form B. Make sure the file contains course codes such as CSC201 or CSC-201 and course titles.")
    return courses
